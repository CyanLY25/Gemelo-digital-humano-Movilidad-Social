import io
import pandas as pd
from fastapi import APIRouter, HTTPException, Response
from fastapi.responses import StreamingResponse
from backend.models.schemas import ExportReportRequest
from backend.services.hdt_abm_engine import run_hdt_simulation, PRESET_POLICIES

router = APIRouter(prefix="/api/export", tags=["Export Reports"])

@router.post("/report")
async def export_report_endpoint(request: ExportReportRequest):
    """Genera y retorna reportes ejecutivos en formato PDF, Word (.docx) o Excel (.xlsx)."""
    try:
        sim_data = run_hdt_simulation(country_id=request.country_id, policy=PRESET_POLICIES["base"], num_agents=500)
        
        if request.format_type == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors
            
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
            styles = getSampleStyleSheet()
            story = []
            
            story.append(Paragraph("<b>GEMELO DIGITAL HUMAN SOCIAL — REPORTE EJECUTIVO v3.11</b>", styles['Title']))
            story.append(Spacer(1, 12))
            story.append(Paragraph(f"<b>País / Calibración:</b> {sim_data.country.name}", styles['Normal']))
            story.append(Paragraph(f"<b>Coeficiente de Gini Proyectado:</b> {sim_data.estimatedGini:.3f}", styles['Normal']))
            story.append(Paragraph(f"<b>Elasticidad Intergeneracional (IGE):</b> {sim_data.estimatedIGE:.3f}", styles['Normal']))
            story.append(Paragraph(f"<b>Tasa de Ascenso Q1 ➔ Q5:</b> {sim_data.q1ToQ5MobilityRate * 100:.1f}%", styles['Normal']))
            story.append(Spacer(1, 15))
            
            story.append(Paragraph("<b>Matriz de Transición Intergeneracional (5x5 Quintiles)</b>", styles['Heading2']))
            matrix_data = [["Origen \\ Destino", "Q1", "Q2", "Q3", "Q4", "Q5"]]
            labels = ["Q1", "Q2", "Q3", "Q4", "Q5"]
            for idx, row in enumerate(sim_data.transitionMatrix):
                matrix_data.append([labels[idx]] + [f"{val*100:.1f}%" for val in row])
                
            t = Table(matrix_data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.navy),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('GRID', (0,0), (-1,-1), 1, colors.black),
            ]))
            story.append(t)
            
            doc.build(story)
            buffer.seek(0)
            return StreamingResponse(
                buffer,
                media_type="application/pdf",
                headers={"Content-Disposition": f"attachment; filename=Reporte_Ejecutivo_{request.country_id}.pdf"}
            )

        elif request.format_type == "docx":
            import docx
            from docx.shared import Inches, Pt, RGBColor
            
            doc = docx.Document()
            doc.add_heading('GEMELO DIGITAL HUMAN SOCIAL — REPORTE EJECUTIVO v3.11', 0)
            doc.add_paragraph(f"Calibración Regional: {sim_data.country.name}")
            doc.add_paragraph(f"Coeficiente de Gini: {sim_data.estimatedGini:.3f}")
            doc.add_paragraph(f"Elasticidad IGE: {sim_data.estimatedIGE:.3f}")
            doc.add_paragraph(f"Movilidad Q1 -> Q5: {sim_data.q1ToQ5MobilityRate*100:.1f}%")
            
            doc.add_heading('Matriz de Transición Intergeneracional 5x5', level=1)
            table = doc.add_table(rows=6, cols=6)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Origen / Destino'
            for k in range(5):
                hdr_cells[k+1].text = f'Q{k+1}'
                
            labels = ["Q1", "Q2", "Q3", "Q4", "Q5"]
            for r_idx, row in enumerate(sim_data.transitionMatrix):
                row_cells = table.rows[r_idx+1].cells
                row_cells[0].text = labels[r_idx]
                for c_idx, val in enumerate(row):
                    row_cells[c_idx+1].text = f"{val*100:.1f}%"
                    
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return StreamingResponse(
                buffer,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename=Reporte_Ejecutivo_{request.country_id}.docx"}
            )

        elif request.format_type == "xlsx":
            import openpyxl
            
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                # Sheet 1: Resumen Ejecutivo
                df_summary = pd.DataFrame([{
                    "Pais": sim_data.country.name,
                    "Gini": sim_data.estimatedGini,
                    "IGE": sim_data.estimatedIGE,
                    "Mobility_Q1_Q5": sim_data.q1ToQ5MobilityRate,
                    "Pobreza_Multidimensional": sim_data.multidimensionalPovertyRate,
                }])
                df_summary.to_excel(writer, sheet_name="Resumen Ejecutivo", index=False)
                
                # Sheet 2: Microdatos Agentes
                agent_rows = []
                for a in sim_data.agents:
                    agent_rows.append({
                        "ID": a.id, "Gender": a.gender, "OriginQ": a.originQuintile,
                        "FinalQ": a.finalQuintile, "FinalWage": a.finalWage, "Degree": a.highestDegree,
                        "Cognitive": a.capital.cognitive, "Human": a.capital.human,
                        "Social": a.capital.social, "Psychological": a.capital.psychological,
                        "Structural": a.capital.structural
                    })
                pd.DataFrame(agent_rows).to_excel(writer, sheet_name="Microdatos Agentes", index=False)
                
                # Sheet 3: Matriz Transicion
                pd.DataFrame(sim_data.transitionMatrix, columns=["Q1","Q2","Q3","Q4","Q5"], index=["Q1","Q2","Q3","Q4","Q5"]).to_excel(writer, sheet_name="Matriz Transicion 5x5")
                
                # Sheet 4: Auditoria Etica
                df_ethics = pd.DataFrame([{
                    "Certificación": "IEEE 7000-2021",
                    "Demographic_Parity": sim_data.fairnessMetrics.demographicParityGender,
                    "Disparate_Impact_4_5_Rule": sim_data.fairnessMetrics.disparateImpactMinority,
                    "Equalized_Odds_Diff": sim_data.fairnessMetrics.equalizedOddsDifference
                }])
                df_ethics.to_excel(writer, sheet_name="Auditoria Etica", index=False)
                
            buffer.seek(0)
            return StreamingResponse(
                buffer,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename=Reporte_Ejecutivo_{request.country_id}.xlsx"}
            )
            
        else:
            raise HTTPException(status_code=400, detail="Formato no soportado. Use 'pdf', 'docx' o 'xlsx'.")

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar reporte: {str(e)}")
